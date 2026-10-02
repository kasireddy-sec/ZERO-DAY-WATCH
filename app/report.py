from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

OUT = Path("reports/zero_day_intelligence.xlsx")

HEADERS = [
"Record_ID","First_Seen","Last_Seen","State","Confidence","Title","Vendor","Product",
"Affected_Versions","CVE","Alternate_IDs","Exploitation","Patch_Status","Mitigation",
"Fixed_Versions","Impact","Recommended_Action","Primary_Source","Source_URL",
"Source_Count","Alert_Status","Alert_Key","Last_Alert_UTC"
]

def write_excel(vulns):
    OUT.parent.mkdir(exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Zero_Day_Register"
    ws.append(HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.font = Font(bold=True, color="FFFFFF")
        c.alignment = Alignment(wrap_text=True)
    for rid, v in sorted(vulns.items()):
        ws.append([v.get(h,"") for h in HEADERS])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        letter = col[0].column_letter
        ws.column_dimensions[letter].width = min(max(max(len(str(x.value or "")) for x in col)+2, 12), 45)
    wb.save(OUT)
